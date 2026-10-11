import io
from pathlib import Path
from pypdf import PdfReader
from docx import Document

from app.rag.chunker import chunk_text
from app.rag.vector_store import VectorStore

from app.core.logger import get_logger
from app.core.config import settings


logger = get_logger(__name__)


def extract_text_from_bytes(content: bytes, filename: str) -> str:
    """Extract text from raw file bytes based on file extension."""
    ext = Path(filename).suffix.lower()

    if ext in {".txt", ".md", ".csv"}:
        try:
            return content.decode("utf-8")
        except UnicodeDecodeError:
            return content.decode("latin-1", errors="replace")

    elif ext == ".pdf":
        try:
            reader = PdfReader(io.BytesIO(content))
            text_parts = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text_parts.append(extracted.strip())
            return "\n\n".join(text_parts)
        except Exception as e:
            raise ValueError(f"Failed to extract text from PDF: {str(e)}")

    elif ext == ".docx":
        try:
            doc = Document(io.BytesIO(content))
            text_parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        text_parts.append(row_text)
            return "\n\n".join(text_parts)
        except Exception as e:
            raise ValueError(f"Failed to extract text from DOCX: {str(e)}")

    else:
        raise ValueError(f"Unsupported file format '{ext}'. Allowed: .pdf, .docx, .txt, .md, .csv")


def load_document(file_path: str) -> str:
    """Read document from disk and extract its text."""
    logger.info("Loading Document: %s", file_path)
    path = Path(file_path)
    content = path.read_bytes()
    return extract_text_from_bytes(content, path.name)

def ingest_document(file_path: str):
    """ Document - text - chunks -embedding -vector store"""

    text = load_document(file_path)

    chunks = chunk_text( text, chunk_size=300, chunk_overlap=50 )

    vector_store = VectorStore()
    vector_store.build(chunks)

    vector_store.save(
        settings.storage_dir
    )
    logger.info("Document ingestion complete.")
    return vector_store


def ingest_text(text: str, vector_store: VectorStore):
    """Chunk text, add to vector store, and save."""
    chunks = chunk_text(text, chunk_size=settings.chunk_size, chunk_overlap=settings.chunk_overlap)
    if not chunks:
        return 0

    count = vector_store.add_chunks(chunks)
    vector_store.save(settings.storage_dir)
    logger.info("Ingested %d new chunks into vector store.", count)
    return count

